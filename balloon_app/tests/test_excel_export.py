"""Tests for the AS9102 Form 3 Excel exporter."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from balloon_app.data_model import Balloon, BalloonSource, Drawing, Project, ReviewStatus
from balloon_app.excel_export import COLUMNS, FORM_NUMBER, INSPECTION_SHEET_NAME, export_excel


@pytest.fixture
def sample_project_and_balloons():
    project = Project(
        name="Test Project", part_number="PN-123", part_name="Camera Housing",
        revision="A", customer="Acme", unit="in", serial_lot_number="N/A",
        fai_report="N/A", po_number="PO-1", mfg_wo="WO-1",
    )
    drawing = Drawing(project_id=project.id, file_name="part.pdf", original_path="C:/drawings/part.pdf", page_count=1)
    project.drawings.append(drawing)

    balloons = [
        Balloon(
            number=1, drawing_id=drawing.id, page_number=0, x=10, y=10,
            char_type="linear_dimension", raw_text="50.00 ±0.05", note="Length",
            nominal=50.0, tol_plus=0.05, tol_minus=0.05,
            lower_limit=49.95, upper_limit=50.05, inspection_method="Caliper",
            source=BalloonSource.AUTO.value, status=ReviewStatus.ACCEPTED.value, confidence=0.9,
        ),
        Balloon(
            number=2, drawing_id=drawing.id, page_number=0, x=20, y=20,
            char_type="diameter", raw_text="Ø25", note="Diameter", nominal=25.0,
            source=BalloonSource.AUTO.value, status=ReviewStatus.PENDING.value, confidence=0.4,
        ),
        Balloon(
            number=3, drawing_id=drawing.id, page_number=0, x=30, y=30,
            char_type="note", raw_text="REJECTED NOTE", note="Note 1",
            source=BalloonSource.AUTO.value, status=ReviewStatus.REJECTED.value, confidence=0.2,
        ),
        Balloon(
            number=4, drawing_id=drawing.id, page_number=0, x=40, y=40,
            char_type="surface_finish", raw_text="Ra 1.6", surface_finish="Ra 1.6",
            note="Surface Finish", inspection_method="Visual",
            source=BalloonSource.MANUAL.value, status=ReviewStatus.ACCEPTED.value, confidence=1.0,
        ),
    ]
    project.balloons.extend(balloons)
    return project, drawing, balloons


def test_headers_match_spec(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=False)

    wb = openpyxl.load_workbook(out_path)
    assert INSPECTION_SHEET_NAME in wb.sheetnames

    ws = wb[INSPECTION_SHEET_NAME]
    header_row = [ws.cell(row=8, column=i + 1).value for i in range(len(COLUMNS))]
    expected = [title for _letter, title in COLUMNS]
    assert header_row == expected


def test_title_block_and_info_grid(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=False)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[INSPECTION_SHEET_NAME]
    assert ws["A2"].value == "First Article Inspection Report"
    assert "Form 3" in ws["A3"].value
    assert ws["B4"].value == "PN-123"
    assert ws["D4"].value == "Camera Housing"
    assert ws["B5"].value == "A"


def test_default_export_excludes_pending_and_rejected(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=False)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[INSPECTION_SHEET_NAME]
    # 2 rows expected: balloon #1 (accepted) and #4 (manual); #2 pending and #3 rejected excluded.
    char_numbers = [ws.cell(row=r, column=1).value for r in range(9, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert char_numbers == [1, 4]


def test_include_pending_option(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection_pending.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=True)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[INSPECTION_SHEET_NAME]
    char_numbers = [ws.cell(row=r, column=1).value for r in range(9, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    # Rejected (#3) must never be included, even with include_pending=True.
    assert 3 not in char_numbers
    assert 2 in char_numbers


def test_requirement_and_limits_formatting(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=False)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[INSPECTION_SHEET_NAME]
    # Balloon #1 -> row 9: Requirement = nominal, Upper/Lower = signed tolerance deltas.
    assert ws["A9"].value == 1
    assert ws["B9"].value == "Length"
    assert ws["C9"].value == "50"
    assert ws["D9"].value == "in"
    assert ws["F9"].value == "0.05"
    assert ws["G9"].value == "-0.05"
    assert ws["I9"].value == "Caliper"


def test_form_footer_present(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=False)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[INSPECTION_SHEET_NAME]
    footer_values = [ws.cell(row=r, column=9).value for r in range(1, ws.max_row + 1)]
    assert any(v and FORM_NUMBER in v for v in footer_values)


def test_header_frozen_and_autofilter(tmp_path: Path, sample_project_and_balloons):
    project, drawing, balloons = sample_project_and_balloons
    out_path = tmp_path / "inspection.xlsx"
    export_excel(project, drawing, balloons, out_path, include_pending=False)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[INSPECTION_SHEET_NAME]
    assert ws.freeze_panes == "A9"
    assert ws.auto_filter.ref is not None
